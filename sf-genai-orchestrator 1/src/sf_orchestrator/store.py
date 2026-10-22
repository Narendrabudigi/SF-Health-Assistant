"""Everything that reads from or writes to Supabase, apart from the run records of the dashboard.

Input:  the ML notebooks' insights (one table, INSIGHTS_TABLE) and the company's own standards file (an xlsx in a Storage bucket).
Output: the finished reports (one table, REPORTS_TABLE).

The industry standards and the custom standards use the same code with different tables (see build_store).
"""
from __future__ import annotations

import io
import json
import logging
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
from postgrest.exceptions import APIError
from supabase import Client, create_client

from .config import Settings

log = logging.getLogger(__name__)

# ── what moves between the program and the database ──
def metric_key(name: str) -> str:
    """The file-safe key of a metric, used in report file names: "Full & Final Settlement Timeliness" becomes
    "FullAndFinalSettlementTimeliness"; a PascalCase name such as "TimeToHire" stays as it is. The website uses the same rule
    to find a metric's report."""
    words = re.findall(r"[A-Za-z0-9]+", name.replace("&", " and "))
    return "".join(w[:1].upper() + w[1:] for w in words) or "Metric"

@dataclass
class InsightRecord:
    run_date: str  # YYYY-MM-DD
    module: str  # RCM | ONB2 | OFB2 | EC | ECP
    metric_name: str
    payload: Any  # raw ML JSON; validated by the orchestrator
    id: str | None = None  # id of the insight row; stored on the report so the website can find the insight

@dataclass
class ReportRecord:
    run_date: str
    module: str
    metric_name: str  # the metric's file-safe key (see metric_key); the readable name is report["metricName"]
    input_hash: str
    prompt_sha: str
    status: str  # 'succeeded' | 'failed'
    generated_at: str
    insight_id: str | None = None
    provider: str | None = None
    model: str | None = None
    report: dict[str, Any] | None = None
    error: str | None = None
    llm_meta: dict[str, Any] = field(default_factory=dict)  # fallbacks, repairs, tokens, latency

class Store(Protocol):
    """The orchestrator depends only on this, so tests can swap in memory."""

    def list_insights(self, run_date: str, module: str | None = None, metric: str | None = None) -> list[InsightRecord]: ...

    def find_report(self, run_date: str, module: str, metric_name: str) -> ReportRecord | None: ...

    def save_report(self, record: ReportRecord) -> None:
        """Idempotent: one record per (run_date, module, metric_name)."""
        ...

# ── input: the insights of the ML notebooks ──


class TableInsights:
    """The ML notebooks' output, read from ONE Supabase table (INSIGHTS_TABLE).

    One row is one metric of one notebook run: `Module_Name`, `Metric_Name`, the insight itself in `JSON_Result`, and
    `created_at` / `execution_id`, which the notebook writes. Nothing here is ever written.

    Which rows belong to a run date: the rows CREATED on that day (UTC). If a metric has more than one row on that day, the
    newest (highest execution_id) is used. The table name is case-sensitive; the column names are fixed below."""

    COLUMNS = ("execution_id", "created_at", "Module_Name", "Metric_Name", "JSON_Result")
    PAGE = 1000  # Supabase returns at most 1000 rows per request

    def __init__(self, url: str, service_role_key: str, table: str, client: Client | None = None):
        self._db: Client = client or create_client(url, service_role_key)
        self._table = table

    def check(self, write_test: bool = False) -> None:
        """Fail BEFORE any model is called if the table cannot be read. Nothing is written, so write_test changes nothing."""
        try:
            self._db.table(self._table).select(",".join(self.COLUMNS)).limit(1).execute()
        except APIError as err:
            raise RuntimeError(self._explain(err)) from None

    def _explain(self, err: APIError) -> str:
        code = str(getattr(err, "code", "") or "")
        message = str(getattr(err, "message", "") or err)
        low = message.lower()
        if code == "PGRST205" or "find the table" in low:
            hint = f"The table '{self._table}' was not found. Check INSIGHTS_TABLE: table names are case-sensitive."
        elif code in ("42703", "PGRST204") or ("column" in low and ("does not exist" in low or "find the" in low)):
            hint = f"The table must have these columns: {', '.join(self.COLUMNS)}."
        elif code == "42501" or "permission denied" in low:
            hint = "Permission denied. Use the secret / service_role key, not the public key, in SUPABASE_SERVICE_ROLE_KEY."
        elif "jwt" in low or "api key" in low or code.startswith("PGRST30"):
            hint = "The key was not accepted. Check SUPABASE_SERVICE_ROLE_KEY and SUPABASE_URL."
        else:
            hint = "Check SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY and INSIGHTS_TABLE."
        return f"Cannot read the insights table '{self._table}' ({code or 'error'}: {message}). {hint}"

    def _rows_created_on(self, run_date: str) -> list[dict[str, Any]]:
        start = date.fromisoformat(run_date)
        rows: list[dict[str, Any]] = []
        while True:
            page = (
                self._db.table(self._table)
                .select(",".join(self.COLUMNS))
                .gte("created_at", f"{start.isoformat()}T00:00:00Z")
                .lt("created_at", f"{(start + timedelta(days=1)).isoformat()}T00:00:00Z")
                .order("execution_id")
                .range(len(rows), len(rows) + self.PAGE - 1)
                .execute()
                .data
                or []
            )
            rows.extend(page)
            if len(page) < self.PAGE:
                return rows

    def list_insights(self, run_date: str, module: str | None = None, metric: str | None = None) -> list[InsightRecord]:
        rows = self._rows_created_on(run_date)
        newest: dict[tuple[str, str], tuple[tuple[int, str], InsightRecord]] = {}
        duplicates = 0
        for row in rows:
            payload = row.get("JSON_Result")
            if isinstance(payload, str):  # stored as text instead of jsonb
                try:
                    payload = json.loads(payload)
                except json.JSONDecodeError:
                    payload = None
            if not isinstance(payload, dict):
                log.error("%s row %s: JSON_Result is not a JSON object", self._table, row.get("execution_id"))
                payload = None
            mod = str(row.get("Module_Name") or (payload or {}).get("module") or "UNKNOWN")
            shown = (payload or {}).get("metricName") if isinstance((payload or {}).get("metricName"), str) else None
            name = metric_key(shown or str(row.get("Metric_Name") or "Metric"))
            if payload is not None:
                payload = {"module": mod, **payload}  # a module written in the JSON itself wins and is checked later
            if (module and mod.upper() != module.upper()) or (metric and metric not in (name, shown)):
                continue
            record = InsightRecord(run_date=run_date, module=mod, metric_name=name, payload=payload, id=f"{self._table}#{row.get('execution_id')}")
            rank = (int(row.get("execution_id") or 0), str(row.get("created_at") or ""))
            known = newest.get((mod.upper(), name))
            if known is not None:
                duplicates += 1
            if known is None or rank > known[0]:
                newest[(mod.upper(), name)] = (rank, record)
        records = [record for _, record in newest.values()]
        if not rows:
            self._explain_empty(run_date)
        if duplicates:
            log.info("Used the newest of several rows for the same metric: %d older row(s) ignored", duplicates)
        log.info("Read %d insight row(s) for %s from table %s", len(records), run_date, self._table)
        return records

    def _explain_empty(self, run_date: str) -> None:
        """Nothing to process: say WHY (empty table or another day), because 'read 0 rows' alone does not tell."""
        try:
            newest = self._db.table(self._table).select("created_at").order("created_at", desc=True).limit(1).execute().data
        except Exception:  # noqa: BLE001  (a diagnostic must never break the run)
            return
        if not newest:
            log.warning("The insights table '%s' is empty. Has the notebook run and written its insights?", self._table)
        else:
            log.warning(
                "The insights table '%s' has no row created on %s (UTC). The newest row was created %s. "
                "Run the notebooks for this date, or use --run-date with the date of the rows.",
                self._table, run_date, newest[0].get("created_at"),
            )

# ── output: the reports ──


class TableReports:
    """The reports as rows of ONE Supabase table (see sql/llm_reports.sql): ONE row per (module, metric_name), the newest.
    It is written with an upsert, so a new run REPLACES the row of the metric instead of adding a row for a new day; `run_date` says which
    day's insight the row was made from. A failed run never replaces a good row (the orchestrator checks that before it saves).

    The whole report goes into the `report` column exactly as produced, and no column is derived from what is inside it, so a
    new report version never needs a new column. The two facts the program needs back, the fingerprint that decides whether a
    report is still current (inputHash) and the id of the insight row (insightId), are kept inside `llm_meta`."""

    COLUMNS = ("run_date", "module", "metric_name", "status", "provider", "model", "report", "error", "llm_meta", "prompt_sha", "generated_at")
    KEY = "module,metric_name"
    _PROBE = "_preflight"

    def __init__(self, url: str, service_role_key: str, table: str, client: Client | None = None):
        self._db: Client = client or create_client(url, service_role_key)
        self._table = table

    # ── the check before any model is called ──

    def check(self, write_test: bool = False) -> None:
        """Fail BEFORE any model is called (and paid for) if the table cannot be used: wrong name, missing column, or a
        missing rule. With write_test the real save is rehearsed on a throw-away row (written twice, then removed): that
        proves the table accepts a write and REPLACES the row of the same module and metric_name, whatever its run_date."""
        try:
            self._db.table(self._table).select(",".join(self.COLUMNS)).limit(1).execute()
            if write_test:
                self._rehearse()
        except APIError as err:
            raise RuntimeError(self._explain(err)) from None

    def _rehearse(self) -> None:
        probe = {
            "run_date": "1970-01-01", "module": self._PROBE, "metric_name": self._PROBE, "status": "failed", "provider": None,
            "model": None, "report": None, "error": "preflight check", "llm_meta": {}, "prompt_sha": "preflight",
            "generated_at": "1970-01-01T00:00:00+00:00",
        }
        try:
            for day in ("1970-01-01", "1970-01-02"):  # two different days: the second write must replace the first
                self._db.table(self._table).upsert({**probe, "run_date": day}, on_conflict=self.KEY).execute()
            rows = self._db.table(self._table).select("run_date").eq("module", self._PROBE).execute().data
            if len(rows) != 1:
                raise RuntimeError(
                    f"The table '{self._table}' kept {len(rows)} rows for the same module and metric_name instead of replacing "
                    "the row. Run sql/llm_reports.sql in the Supabase SQL Editor: it adds the rule."
                )
        finally:
            try:
                self._db.table(self._table).delete().eq("module", self._PROBE).execute()
            except APIError:
                log.warning("The throw-away row with module '%s' could not be removed from '%s'; delete it by hand.", self._PROBE, self._table)

    def _explain(self, err: APIError) -> str:
        code = str(getattr(err, "code", "") or "")
        message = str(getattr(err, "message", "") or err)
        low = message.lower()
        name = self._table
        if code == "PGRST205" or "find the table" in low:
            hint = f"The table '{name}' was not found. Create it with sql/llm_reports.sql and check REPORTS_TABLE: table names are case-sensitive."
        elif code in ("42703", "PGRST204") or ("column" in low and ("does not exist" in low or "find the" in low)):
            hint = f"The table must have these columns: id, {', '.join(self.COLUMNS)}. Compare it with sql/llm_reports.sql."
        elif code == "42P10" or "no unique or exclusion constraint" in low:
            hint = (
                "The table has no rule 'one row per module and metric_name', so a new run cannot replace a row. "
                "Run sql/llm_reports.sql in the Supabase SQL Editor: it adds the rule."
            )
        elif code in ("23514", "23502", "23505"):
            hint = "A rule of the table rejected the test row. Compare the table with sql/llm_reports.sql."
        elif code in ("22P02", "22007", "22008") or "invalid input syntax" in low:
            hint = "A column has another type than expected: run_date date, generated_at timestamptz, report and llm_meta jsonb."
        elif code == "42501" or "permission denied" in low:
            hint = "Permission denied. Use the secret / service_role key, not the public key, in SUPABASE_SERVICE_ROLE_KEY."
        elif "jwt" in low or "api key" in low or code.startswith("PGRST30"):
            hint = "The key was not accepted. Check SUPABASE_SERVICE_ROLE_KEY and SUPABASE_URL."
        else:
            hint = "Check SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY and that the table exists."
        return f"Cannot use the reports table '{name}' ({code or 'error'}: {message}). {hint}"

    # ── reading and writing ──

    def find_report(self, run_date: str, module: str, metric_name: str) -> ReportRecord | None:
        """The row of the metric, whatever day it was made on (run_date is not used: there is one row per metric)."""
        rows = (
            self._db.table(self._table)
            .select(",".join(self.COLUMNS))
            .eq("module", module)
            .eq("metric_name", metric_name)
            .limit(1)
            .execute()
            .data
        )
        if not rows:
            return None
        r = rows[0]
        meta = dict(r.get("llm_meta") or {})
        return ReportRecord(
            run_date=r["run_date"],
            module=r["module"],
            metric_name=r["metric_name"],
            input_hash=meta.pop("inputHash", "") or "",
            prompt_sha=r["prompt_sha"],
            status=r["status"],
            generated_at=r["generated_at"],
            insight_id=meta.pop("insightId", None),
            provider=r.get("provider"),
            model=r.get("model"),
            report=r.get("report"),
            error=r.get("error"),
            llm_meta=meta,
        )

    def save_report(self, r: ReportRecord) -> None:
        self._db.table(self._table).upsert(
            {
                "run_date": r.run_date,
                "module": r.module,
                "metric_name": r.metric_name,
                "status": r.status,
                "provider": r.provider,
                "model": r.model,
                "report": r.report,  # the complete JSON, as produced
                "error": r.error,
                "llm_meta": {**r.llm_meta, "inputHash": r.input_hash, "insightId": r.insight_id},
                "prompt_sha": r.prompt_sha,
                "generated_at": r.generated_at,
            },
            on_conflict=self.KEY,
        ).execute()

class CombinedStore:
    """Insights from one place, reports to another: for example insights from a table and reports to another table."""

    def __init__(self, insights: Store, reports: Store):
        self._insights, self._reports = insights, reports

    def check(self, write_test: bool = False) -> None:
        for part in (self._insights, self._reports):
            if hasattr(part, "check"):
                part.check(write_test=write_test)

    def list_insights(self, run_date: str, module: str | None = None, metric: str | None = None) -> list[InsightRecord]:
        return self._insights.list_insights(run_date, module, metric)

    def find_report(self, run_date: str, module: str, metric_name: str) -> ReportRecord | None:
        return self._reports.find_report(run_date, module, metric_name)

    def save_report(self, record: ReportRecord) -> None:
        self._reports.save_report(record)

# ── which tables a run uses ──


def build_store(settings: Settings, custom: bool = False) -> CombinedStore:
    """The tables a run uses: the industry standards read INSIGHTS_TABLE and write REPORTS_TABLE; the company's own standards
    read CUSTOM_INSIGHTS_TABLE and write CUSTOM_REPORTS_TABLE. load_settings has already checked that the names are set."""
    insights_table, reports_table = (
        (settings.custom_insights_table, settings.custom_reports_table) if custom else (settings.insights_table, settings.reports_table)
    )
    assert settings.supabase_url and settings.supabase_service_role_key and insights_table and reports_table
    key = settings.supabase_service_role_key.get_secret_value()
    label = "Custom standards: " if custom else ""
    log.info("%sInsights: read from the table '%s' (the rows created on the run date, UTC)", label, insights_table)
    log.info("%sReports: saved as rows of the table '%s'", label, reports_table)
    return CombinedStore(TableInsights(settings.supabase_url, key, insights_table), TableReports(settings.supabase_url, key, reports_table))

# ── input: the company's own standards file ──
# The company's custom standards: an xlsx file the client uploads to a Storage bucket.
# The file has one row per metric with the columns `Module`, `Metric Name`, `Industry Benchmark` and `Unit`. The custom part of the
# pipeline runs only for what the client has filled in: a metric with an empty benchmark cell is skipped, a module with no
# value at all is skipped, and a file with no value at all (or no file) means that nothing of the custom part runs.
# 
# This file only READS the standards to decide that. Turning the values into insights is done by the custom notebooks, which read the
# same file (CUSTOM_STANDARDS_FILE_PATH).

MODULES = ("RCM", "ONB2", "OFB2", "EC", "ECP")
_EMPTY_WORDS = {"", "nan", "n/a", "na", "null", "none", "-", "--"}


class StandardsError(Exception):
    """The standards file exists but cannot be used (wrong columns, not an xlsx, or the bucket cannot be reached)."""


@dataclass(frozen=True)
class Standard:
    module: str
    metric: str
    has_value: bool
    stored_as_date: bool = False  # Excel turned a number such as 5 into the date 1900-01-05: the cell has the wrong format
    number_behind_date: float | None = None


@dataclass
class StandardsCheck:
    """What the file says. `found` is False when the file is not in the bucket (the client has not uploaded standards yet)."""

    found: bool
    standards: list[Standard] = field(default_factory=list)

    @property
    def modules_with_values(self) -> set[str]:
        return {s.module for s in self.standards if s.has_value}

    @property
    def blank(self) -> list[Standard]:
        return [s for s in self.standards if not s.has_value]

    @property
    def ready(self) -> bool:
        return self.found and bool(self.modules_with_values)


def _has_value(cell: object) -> bool:
    if cell is None:
        return False
    if isinstance(cell, float) and math.isnan(cell):
        return False
    if isinstance(cell, str):
        return cell.strip().lower() not in _EMPTY_WORDS
    return True  # a number, a date or anything else the client typed counts as a value


def _header_index(row: tuple) -> dict[str, int] | None:
    names = {str(v).strip().lower(): i for i, v in enumerate(row) if v is not None}
    wanted = ("module", "metric name", "industry benchmark")
    return {k: names[k] for k in wanted} if all(k in names for k in wanted) else None


def parse_standards(content: bytes) -> list[Standard]:
    """Every metric row of the file. A sheet is used when its first rows hold the three column names (in any order, any capitals)."""
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as err:  # not an xlsx, damaged, password protected ...
        raise StandardsError(f"The custom standards file cannot be opened as an Excel (.xlsx) file: {type(err).__name__}") from None
    standards: list[Standard] = []
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = None
            for row in rows:
                header = _header_index(row)
                if header:
                    break
            if not header:
                continue
            for row in rows:
                get = lambda key: row[header[key]] if header[key] < len(row) else None  # noqa: E731
                module, metric = get("module"), get("metric name")
                if module is None or metric is None or not str(module).strip() or not str(metric).strip():
                    continue
                benchmark = get("industry benchmark")
                is_date = isinstance(benchmark, (datetime, date))
                standards.append(
                    Standard(
                        module=str(module).strip().upper(),
                        metric=re.sub(r"\s+", " ", str(metric)).strip(),
                        has_value=_has_value(benchmark),
                        stored_as_date=is_date,
                        number_behind_date=float(to_excel(benchmark)) if is_date else None,
                    )
                )
    finally:
        book.close()
    if not standards:
        raise StandardsError(
            "The custom standards file has no sheet with the columns Module, Metric Name and Industry Benchmark, or it has no rows"
        )
    return standards


def fetch_standards(url: str, service_role_key: str, bucket: str, path: str, client=None) -> StandardsCheck:
    """Read the file from the bucket. A file that is not there is not an error: the client has simply not uploaded any standards."""
    db = client or create_client(url, service_role_key)
    try:
        content = db.storage.from_(bucket).download(path)
    except Exception as err:  # noqa: BLE001  (the storage library raises its own error types)
        text = str(err).lower()
        if "not found" in text or "404" in text:
            return StandardsCheck(found=False)
        raise StandardsError(f"The custom standards file '{path}' could not be read from the bucket '{bucket}': {type(err).__name__}: {str(err)[:200]}") from None
    return StandardsCheck(found=True, standards=parse_standards(content))


def notebook_module(notebook: Path) -> str | None:
    """The module a notebook belongs to: its file name starts with the module code, for example `EC Notebook1.py` or `rcm_time_to_hire.py`.
    None when the name does not start with one (such a notebook is run whenever anything of the custom part runs)."""
    first = re.split(r"[ _.\-]", notebook.stem, maxsplit=1)[0].upper()
    return first if first in MODULES else None


def describe(check: StandardsCheck) -> list[str]:
    """Log lines about what the file contains: which metrics are skipped, and cells that look wrong."""
    lines = []
    for s in check.standards:
        if s.stored_as_date:
            lines.append(
                f"{s.module} / {s.metric}: the benchmark cell is formatted as a DATE; it holds the number {s.number_behind_date:g}. "
                f"Format the cell as a number (or General) so that the notebooks read it correctly"
            )
    skipped = [f"{s.module} / {s.metric}" for s in check.blank]
    if skipped:
        lines.append(f"No benchmark value for {len(skipped)} metric(s), they are skipped: {'; '.join(skipped)}")
    return lines