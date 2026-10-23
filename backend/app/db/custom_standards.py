"""Read and replace the custom standards workbook in Supabase Storage."""
from __future__ import annotations

from io import BytesIO
import re
import threading
import time
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.config import settings
from app.db.mock_data import DEFAULT_MODULES
from app.db.supabase_client import supabase_client, supabase_storage

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_WORKBOOK_BYTES = 15 * 1024 * 1024
CUSTOM_STANDARDS_CACHE_TTL_SECONDS = 30
_custom_standards_cache: tuple[float, dict[str, Any]] | None = None
_custom_standards_cache_lock = threading.Lock()


def invalidate_custom_standards_cache() -> None:
    """Drop the parsed workbook cache after a standards upload."""
    global _custom_standards_cache
    with _custom_standards_cache_lock:
        _custom_standards_cache = None


def _alpha(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _module(value: Any) -> tuple[str, str] | None:
    aliases = {
        "ec": ("EC", "ec"),
        "employeecentral": ("EC", "ec"),
        "hr": ("EC", "ec"),
        "rcm": ("RCM", "rcm"),
        "recruitment": ("RCM", "rcm"),
        "recruiting": ("RCM", "rcm"),
        "onb": ("ONB2", "onb"),
        "onb2": ("ONB2", "onb"),
        "onboarding": ("ONB2", "onb"),
        "ofb": ("OFB2", "ofb"),
        "ofb2": ("OFB2", "ofb"),
        "offboarding": ("OFB2", "ofb"),
        "ecp": ("ECP", "ecp"),
        "employeecentralpayroll": ("ECP", "ecp"),
        "payroll": ("ECP", "ecp"),
    }
    return aliases.get(_alpha(value))


def _canonical_metric(module_id: str, metric_name: str) -> tuple[str, str]:
    """Return the exact Supabase Metric_Name plus its readable catalogue label."""
    metric_key = _alpha(metric_name)
    for module in DEFAULT_MODULES:
        if str(module.get("id", "")).lower() != module_id.lower():
            continue
        for benchmark in module.get("benchmarks", []):
            aliases = {
                _alpha(benchmark.get("metric")),
                _alpha(benchmark.get("code")),
                _alpha(benchmark.get("supabaseMetricName")),
            }
            if metric_key in aliases:
                return benchmark.get("supabaseMetricName") or metric_name, benchmark.get("metric") or metric_name
    return metric_name, metric_name


def _standard_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def with_inferred_unit(value: Any, reference: Any) -> str:
    """Add the catalogue unit to a numeric custom standard when its cell omits it."""
    text = _standard_text(value)
    reference_text = _standard_text(reference)
    if not text or not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
        return text
    unit_match = re.search(r"(%|\bpercent(?:age)?s?\b|\bdays?\b)", reference_text, re.IGNORECASE)
    if not unit_match:
        return text
    unit = "%" if unit_match.group(1) == "%" or "percent" in unit_match.group(1).lower() else "Days"
    return f"{text}{unit}" if unit == "%" else f"{text} {unit}"


def parse_custom_standards(content: bytes) -> list[dict[str, str]]:
    """Parse the workbook format consumed by the orchestrator notebooks."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="The standards file must be a readable .xlsx workbook.") from exc

    parsed: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    header_found = False
    try:
        for sheet in workbook.worksheets:
            header_row = None
            header_index = None
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                headers = {_alpha(value): index for index, value in enumerate(row) if value is not None}
                if {"module", "metricname", "industrybenchmark"}.issubset(headers):
                    header_row = row_index
                    header_index = headers
                    header_found = True
                    break
            if header_index is None:
                continue

            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                module_value = row[header_index["module"]] if len(row) > header_index["module"] else None
                metric_value = row[header_index["metricname"]] if len(row) > header_index["metricname"] else None
                standard_value = row[header_index["industrybenchmark"]] if len(row) > header_index["industrybenchmark"] else None
                unit_value = row[header_index["unit"]] if "unit" in header_index and len(row) > header_index["unit"] else None
                metric = str(metric_value or "").strip()
                standard = _standard_text(standard_value)
                unit = str(unit_value or "").strip()
                if standard and unit and re.fullmatch(r"[-+]?\d+(?:\.\d+)?", standard):
                    normalized_unit = unit.lower().strip()
                    if normalized_unit in {"percentage", "percent", "percentage points", "%"}:
                        standard = f"{standard}%"
                    elif normalized_unit in {"day", "days"}:
                        standard = f"{standard} Days"
                    else:
                        standard = f"{standard} {unit}"
                if not metric or not standard:
                    continue
                module = _module(module_value)
                if not module:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Unknown module '{module_value}' for metric '{metric}'. Use EC, RCM, ONB2, OFB2, or ECP.",
                    )
                module_code, module_id = module
                canonical_name, display_name = _canonical_metric(module_id, metric)
                identity = (module_id, _alpha(canonical_name))
                if identity in seen:
                    raise HTTPException(
                        status_code=400,
                        detail=f"The workbook contains duplicate standards for {module_code} / {metric}.",
                    )
                seen.add(identity)
                parsed.append({
                    "module": module_id,
                    "moduleCode": module_code,
                    "metric": canonical_name,
                    "displayMetric": display_name,
                    "standard": standard,
                    "standardValue": standard_value,
                })
    finally:
        workbook.close()

    if not header_found:
        raise HTTPException(
            status_code=400,
            detail="The workbook needs columns named Module, Metric Name, and Industry Benchmark.",
        )
    if not parsed:
        raise HTTPException(status_code=400, detail="The workbook contains no rows with a module, metric, and standard value.")
    return parsed


def standards_response(rows: list[dict[str, str]], *, exists: bool = True) -> dict[str, Any]:
    public_rows = [
        {key: row[key] for key in ("module", "moduleCode", "metric", "standard")}
        for row in rows
    ]
    standards_map = {
        f"{row['module']}:{_alpha(row['metric'])}": row["standard"]
        for row in public_rows
    }
    for row in rows:
        legacy_key = f"{row['module']}:{_alpha(row.get('displayMetric'))}"
        standards_map.setdefault(legacy_key, row["standard"])
    return {
        "exists": exists,
        "bucket": settings.SUPABASE_CUSTOM_STANDARDS_BUCKET_NAME,
        "path": settings.SUPABASE_CUSTOM_STANDARDS_PATH,
        "standards": public_rows,
        "standardsMap": standards_map,
    }


def load_custom_standards(*, force_refresh: bool = False) -> dict[str, Any]:
    """Load parsed standards from a short-lived process cache or Supabase Storage."""
    global _custom_standards_cache
    if not supabase_client:
        raise HTTPException(status_code=503, detail="Supabase is not configured.")
    with _custom_standards_cache_lock:
        if not force_refresh and _custom_standards_cache:
            cached_at, cached_response = _custom_standards_cache
            if time.monotonic() - cached_at < CUSTOM_STANDARDS_CACHE_TTL_SECONDS:
                return cached_response

        try:
            content = supabase_client.storage.from_(settings.SUPABASE_CUSTOM_STANDARDS_BUCKET_NAME).download(
                settings.SUPABASE_CUSTOM_STANDARDS_PATH
            )
            result = standards_response(parse_custom_standards(content))
        except Exception as exc:
            error_text = str(exc).lower()
            if "not found" in error_text or "404" in error_text:
                result = standards_response([], exists=False)
            else:
                raise HTTPException(
                    status_code=502,
                    detail="Could not fetch the custom standards workbook from Supabase Storage.",
                ) from exc

        _custom_standards_cache = (time.monotonic(), result)
        return result


def normalized_workbook(rows: list[dict[str, str]]) -> bytes:
    """Store canonical orchestrator module IDs in the shared workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Custom Standards"
    sheet.append(["Module", "Metric Name", "Industry Benchmark"])
    for row in rows:
        # Keep any normalized unit in the benchmark string when storing the
        # workbook; the parser combines numeric values and Unit into this field.
        sheet.append([row["moduleCode"], row["metric"], row["standard"]])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def custom_standards_template() -> bytes:
    """Create a prefilled upload workbook in the reference workbook's module order."""
    workbook = Workbook()
    standards = workbook.active
    standards.title = "Custom Standards"
    standards.append(["Module", "Metric Name", "Industry Benchmark", "Unit"])
    module_order = ("rcm", "onb", "ofb", "ec", "ecp")
    module_codes = {"rcm": "RCM", "onb": "ONB2", "ofb": "OFB2", "ec": "EC", "ecp": "ECP"}
    for module_id in module_order:
        module = next((item for item in DEFAULT_MODULES if str(item.get("id", "")).lower() == module_id), None)
        if not module:
            continue
        for benchmark in module.get("benchmarks", []):
            standards.append([
                module_codes[module_id],
                benchmark.get("supabaseMetricName") or benchmark.get("metric") or benchmark.get("code"),
                None,
                None,
            ])
    standards.freeze_panes = "A2"
    standards.auto_filter.ref = f"A1:D{standards.max_row}"
    for cell in standards[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="105276")
    for row in standards.iter_rows(min_row=2):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="F3F7FB" if cell.row % 2 == 0 else "FFFFFF")
    for column, width in {"A": 18, "B": 42, "C": 24, "D": 18}.items():
        standards.column_dimensions[column].width = width
    module_validation = DataValidation(type="list", formula1='"EC,RCM,ONB2,OFB2,ECP"', allow_blank=True)
    standards.add_data_validation(module_validation)
    module_validation.add(f"A2:A{standards.max_row}")
    unit_validation = DataValidation(type="list", formula1='"Percentage,Days"', allow_blank=True)
    standards.add_data_validation(unit_validation)
    unit_validation.add(f"D2:D{standards.max_row}")

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Field", "What to enter", "Example"])
    instructions.append(["Module", "Use one of the supported module codes.", "EC"])
    instructions.append(["Metric Name", "All 21 exact Supabase Metric_Name values are prefilled in the reference module order. Keep each name with its module.", "EmployeeDataAccuracy"])
    instructions.append(["Industry Benchmark", "Enter the target number. The Unit column is used to format the value.", "98.5 or 2"])
    instructions.append(["Unit", "Choose Percentage or Days when applicable.", "Percentage or Days"])
    instructions.append(["Percentage standard", "Select Percentage in Unit; the uploaded value will include %.", "98.5 + Percentage"])
    instructions.append(["Days standard", "Select Days in Unit; the uploaded value will include Days.", "2 + Days"])
    instructions.append(["EC metric names", "Use these exact database names.", "EmployeeDataAccuracy; DataSyncErrorRate; WorkflowApprovalCycleTime; AvgTimeToProcessDataChangeRequests; RetroactiveDataChangesFrequency; NewHireEarlyAttritionRate90Day; PositionManagementAccuracy; RegrettableAttritionRate"])
    instructions.append(["RCM metric names", "Use these exact database names.", "TimeToHire; ApplicantToInterviewConversionRate; OfferAcceptanceRate; RequisitionAging; SourceOfHireEffectiveness"])
    instructions.append(["ONB2 metric names", "Use these exact database names.", "OnboardingCycleTime; PreDay1TaskCompletionRate; Day306090MilestoneCompletionRate; ComplianceDocumentationCompletionRate"])
    instructions.append(["OFB2 metric names", "Use these exact database names.", "AccessAssetRevocationTimeliness; OffboardingCycleTime; KnowledgeTransferCompletionRate; FullFinalSettlementTimeliness"])
    instructions.append(["Important", "Keep the Custom Standards sheet headers unchanged. Examples on this sheet are not uploaded as standards.", ""])
    instructions.freeze_panes = "A2"
    for cell in instructions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="105276")
    for column, width in {"A": 24, "B": 86, "C": 40}.items():
        instructions.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload_custom_standards(content: bytes) -> dict[str, Any]:
    if not supabase_client:
        raise HTTPException(status_code=503, detail="Supabase is not configured.")
    rows = parse_custom_standards(content)
    try:
        supabase_client.storage.from_(settings.SUPABASE_CUSTOM_STANDARDS_BUCKET_NAME).upload(
            settings.SUPABASE_CUSTOM_STANDARDS_PATH,
            normalized_workbook(rows),
            {"upsert": "true", "content-type": XLSX_CONTENT_TYPE},
        )
    except Exception as exc:
        raise HTTPException(status_code=502, detail="Could not replace the custom standards workbook in Supabase Storage.") from exc
    supabase_storage.invalidate_cache()
    invalidate_custom_standards_cache()
    return standards_response(rows)
