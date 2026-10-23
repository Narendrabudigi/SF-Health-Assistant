import csv
import csv
import io
import json
import logging
import re
from datetime import date, datetime
from pathlib import PurePosixPath
from typing import Any, Dict, List, Tuple
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response

from app.config import settings
from app.db.supabase_client import supabase_client

router = APIRouter()
logger = logging.getLogger("sf_health.datasets")
SUPPORTED_SUFFIXES = {".xlsx", ".csv"}
NON_DATASET_FILE_MARKERS = {"standard", "standards", "benchmark", "benchmarks", "template", "manifest"}


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _tokens(value: str) -> set:
    separated = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", value or "")
    return {_key(token) for token in re.findall(r"[A-Za-z0-9]+", separated) if len(_key(token)) > 2}


def _list_dataset_objects() -> List[str]:
    if not supabase_client:
        raise HTTPException(status_code=503, detail="Supabase is not configured")

    storage = supabase_client.storage.from_(settings.SUPABASE_DATASET_BUCKET_NAME)
    found: List[str] = []
    visited = set()
    pending = [""]
    while pending:
        folder = pending.pop()
        if folder in visited:
            continue
        visited.add(folder)
        try:
            entries = storage.list(folder) if folder else storage.list()
        except Exception as exc:
            if not folder:
                raise HTTPException(status_code=502, detail=f"Could not list raw dataset bucket '{settings.SUPABASE_DATASET_BUCKET_NAME}': {exc}")
            continue

        for entry in entries or []:
            if not isinstance(entry, dict):
                continue
            name = str(entry.get("name") or "").strip()
            if not name or name.startswith("."):
                continue
            path = f"{folder}/{name}" if folder else name
            if PurePosixPath(name).suffix.lower() in SUPPORTED_SUFFIXES:
                found.append(path)
            elif entry.get("id") is None and not entry.get("metadata"):
                pending.append(path)
    return found


def _load_dataset_manifest() -> Dict[str, Any]:
    manifest_path = settings.SUPABASE_DATASET_MANIFEST_PATH
    if not manifest_path or not supabase_client:
        return {}
    try:
        raw = supabase_client.storage.from_(settings.SUPABASE_DATASET_BUCKET_NAME).download(manifest_path)
        parsed = json.loads(raw.decode("utf-8-sig") if isinstance(raw, bytes) else raw)
        return parsed if isinstance(parsed, dict) else {}
    except Exception as exc:
        # The manifest is optional. Automatic name matching remains available
        # for datasets whose names still identify their module or metric.
        logger.info("Dataset manifest '%s' is unavailable; using automatic matching: %s", manifest_path, exc)
        return {}


def _as_values(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, (str, int, float)):
        return [str(value)]
    if isinstance(value, list):
        return [str(item) for item in value if item is not None]
    return []


def _manifest_dataset(
    manifest: Dict[str, Any], module_id: str, module_name: str,
    metric_name: str, metric_code: str, supabase_metric_name: str = ""
) -> Tuple[str, str]:
    entries = manifest.get("datasets", [])
    if not isinstance(entries, list):
        return "", ""

    module_ids = {_key(module_id), _key(module_name)} - {""}
    metric_ids = {_key(metric_name), _key(metric_code), _key(supabase_metric_name)} - {""}
    exact = []
    defaults = []
    module_has_entries = False

    for entry in entries:
        if not isinstance(entry, dict):
            continue
        entry_modules = {
            _key(value)
            for value in _as_values(entry.get("modules", entry.get("module")))
            + _as_values(entry.get("module_aliases"))
            if _key(value)
        }
        module_matches = not entry_modules or "*" in entry_modules or bool(module_ids & entry_modules)
        if not module_matches:
            continue
        module_has_entries = True

        dataset_path = str(entry.get("path") or entry.get("file") or "").strip().lstrip("/")
        if not dataset_path:
            continue

        metric_values = (
            _as_values(entry.get("metrics"))
            + _as_values(entry.get("metric"))
            + _as_values(entry.get("metric_aliases"))
            + _as_values(entry.get("metric_codes"))
        )
        entry_metrics = {_key(value) for value in metric_values if _key(value)}
        if entry_metrics & metric_ids:
            exact.append(entry)
        elif entry.get("default") is True or "*" in entry_metrics:
            defaults.append(entry)

    selected = exact if exact else defaults
    if len(selected) > 1:
        match_kind = "metric" if exact else "module default"
        raise HTTPException(
            status_code=409,
            detail=f"Dataset manifest has multiple {match_kind} entries matching {module_name or module_id} / {metric_name}.",
        )
    if not selected:
        if module_has_entries:
            raise HTTPException(
                status_code=404,
                detail=(f"No dataset mapping exists for '{metric_name}' in '{module_name or module_id}'. "
                        f"Add the metric to '{settings.SUPABASE_DATASET_MANIFEST_PATH}'."),
            )
        return "", ""

    chosen = selected[0]
    return (
        str(chosen.get("path") or chosen.get("file") or "").strip().lstrip("/"),
        str(chosen.get("sheet") or chosen.get("worksheet") or "").strip(),
    )


def _choose_dataset(
    paths: List[str], module_id: str, metric_name: str,
    metric_code: str = "", module_name: str = "", supabase_metric_name: str = ""
) -> Tuple[str, str]:
    manifest = _load_dataset_manifest()
    manifest_path, sheet_name = _manifest_dataset(
        manifest, module_id, module_name, metric_name, metric_code, supabase_metric_name
    )
    if manifest_path:
        if manifest_path not in paths:
            raise HTTPException(
                status_code=404,
                detail=(f"Dataset manifest points to '{manifest_path}', but that file is not present "
                        f"in bucket '{settings.SUPABASE_DATASET_BUCKET_NAME}'. Update the manifest path."),
            )
        return manifest_path, sheet_name

    # Raw-data buckets can also contain standards workbooks and the optional
    # manifest. Never let those files compete with the actual source dataset.
    dataset_paths = [
        path for path in paths
        if not (_tokens(PurePosixPath(path).stem) & NON_DATASET_FILE_MARKERS)
    ]

    metric_keys = {_key(metric_name), _key(metric_code), _key(supabase_metric_name)} - {""}
    metric_tokens = (_tokens(metric_name) | _tokens(metric_code) | _tokens(supabase_metric_name)) - {"data", "metric", "rate", "average", "total", "percentage", "percent"}
    module_tokens = _tokens(module_id) | _tokens(module_name)
    module_code = _key(module_id)
    module_key = _key(module_id)

    scored = []
    for path in dataset_paths:
        filename_key = _key(PurePosixPath(path).stem)
        path_tokens = _tokens(path)
        module_match = bool(module_tokens & path_tokens) or bool(
            len(module_code) >= 3 and any(token.startswith(module_code) for token in path_tokens)
        )
        score = 0
        if any(key in filename_key for key in metric_keys):
            score += 100
        score += 20 * len(metric_tokens & path_tokens)
        if module_match:
            score += 15
        if module_key == "ec" and any(key in filename_key for key in ("employee", "central", "ec")):
            score += 5
        scored.append((score, path))

    if not scored:
        raise HTTPException(
            status_code=404,
            detail=(f"No source CSV or Excel datasets found in '{settings.SUPABASE_DATASET_BUCKET_NAME}'. "
                    "Only standards, templates, or manifest files were found."),
        )
    if len(scored) == 1:
        return scored[0][1], ""

    scored.sort(reverse=True)
    best_score, best_path = scored[0]
    tied = [path for score, path in scored if score == best_score]
    if best_score < 20 or len(tied) > 1:
        options = [PurePosixPath(path).name for _, path in scored]
        raise HTTPException(
            status_code=404,
            detail=(f"Could not match a dataset to '{metric_name}' in '{module_name or module_id}'. "
                    f"Add a mapping to '{settings.SUPABASE_DATASET_MANIFEST_PATH}' or use descriptive file names. "
                    f"Available datasets: {', '.join(options)}"),
        )
    return best_path, ""


def _download_object(path: str) -> bytes:
    if not supabase_client:
        raise HTTPException(status_code=503, detail="Supabase is not configured")
    try:
        return supabase_client.storage.from_(settings.SUPABASE_DATASET_BUCKET_NAME).download(path)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Could not download dataset from Supabase: {exc}")


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _month_key(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    if not text:
        return ""
    iso_match = re.match(r"^(\d{4})[-/](\d{1,2})", text)
    if iso_match:
        year, month = map(int, iso_match.groups())
        return f"{year:04d}-{month:02d}" if 1 <= month <= 12 else ""
    for date_format in ("%m/%d/%Y", "%d/%m/%Y", "%m-%d-%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], date_format).strftime("%Y-%m")
        except ValueError:
            continue
    return ""


def _select_worksheet(workbook, metric_name: str, metric_code: str, preferred_sheet: str, module_id: str, module_name: str):
    sheet = workbook.active
    if preferred_sheet:
        return next((candidate for candidate in workbook.worksheets if candidate.title == preferred_sheet), None) or sheet
    if len(workbook.worksheets) == 1:
        return sheet

    ignored = {"data", "metric", "rate", "average", "total", "percentage", "percent"}
    metric_tokens = (_tokens(metric_name) | _tokens(metric_code)) - ignored
    module_tokens = _tokens(module_id) | _tokens(module_name)
    sheet_matches = [
        (
            3 * len(metric_tokens & _tokens(candidate.title))
            + len(module_tokens & _tokens(candidate.title))
            + int(
                len(_key(module_id)) >= 3
                and any(token.startswith(_key(module_id)) for token in _tokens(candidate.title))
            ),
            candidate,
        )
        for candidate in workbook.worksheets
    ]
    matched_score, matched_sheet = max(sheet_matches, key=lambda item: item[0])
    if matched_score == 0:
        available = ", ".join(candidate.title for candidate in workbook.worksheets)
        raise HTTPException(
            status_code=404,
            detail=(f"Could not identify a worksheet for '{metric_name}' in "
                    f"'{module_name or module_id}'. Available worksheets: {available}. "
                    f"Add the worksheet to '{settings.SUPABASE_DATASET_MANIFEST_PATH}'."),
        )
    return matched_sheet


def _read_dataset(
    path: str,
    content: bytes,
    offset: int,
    limit: int,
    metric_name: str = "",
    metric_code: str = "",
    preferred_sheet: str = "",
    module_id: str = "",
    module_name: str = "",
    date_column: str = "",
    date_from: str = "",
    date_to: str = "",
) -> Dict[str, Any]:
    suffix = PurePosixPath(path).suffix.lower()
    workbook = None
    if suffix == ".csv":
        text = content.decode("utf-8-sig", errors="replace")
        row_iterator = iter(csv.reader(io.StringIO(text)))
        sheet_name = "CSV"
    else:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = _select_worksheet(workbook, metric_name, metric_code, preferred_sheet, module_id, module_name)
            row_iterator = iter(sheet.iter_rows(values_only=True))
            sheet_name = sheet.title
        except Exception as exc:
            if workbook:
                workbook.close()
            if isinstance(exc, HTTPException):
                raise
            raise HTTPException(status_code=422, detail=f"Could not read Excel dataset: {exc}")

    header = next(row_iterator, None)
    if header is None:
        if workbook:
            workbook.close()
        return {"fileName": PurePosixPath(path).name, "sheetName": sheet_name, "columns": [], "rows": [], "totalRows": 0, "offset": offset, "limit": limit}

    columns = [str(value).strip() if value is not None else f"Column {idx + 1}" for idx, value in enumerate(header)]
    date_columns = [column for column in columns if re.search(r"date|time", column, re.IGNORECASE)]
    date_column_index = next(
        (index for index, column in enumerate(columns) if column.casefold() == date_column.casefold()),
        None,
    ) if date_column else None
    if (date_from or date_to) and date_column_index is None:
        if workbook:
            workbook.close()
        raise HTTPException(status_code=400, detail="Choose a date column before applying a date range.")
    page_rows = []
    total_rows = 0
    for row in row_iterator:
        values = list(row)
        if not any(value is not None and str(value).strip() for value in values):
            continue
        values += [None] * max(0, len(columns) - len(values))
        if date_column_index is not None and (date_from or date_to):
            row_month = _month_key(values[date_column_index])
            if not row_month or (date_from and row_month < date_from) or (date_to and row_month > date_to):
                continue
        if offset <= total_rows < offset + limit:
            page_rows.append({columns[idx]: _cell_value(values[idx]) for idx in range(len(columns))})
        total_rows += 1

    if workbook:
        workbook.close()

    return {
        "fileName": PurePosixPath(path).name,
        "sheetName": sheet_name,
        "columns": columns,
        "dateColumns": date_columns,
        "rows": page_rows,
        "totalRows": total_rows,
        "offset": offset,
        "limit": limit,
    }


@router.get("/preview")
def preview_dataset(
    module_id: str = Query(...),
    metric_name: str = Query(...),
    metric_code: str = Query(""),
    supabase_metric_name: str = Query(""),
    module_name: str = Query(""),
    offset: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=500),
    date_column: str = Query(""),
    date_from: str = Query(""),
    date_to: str = Query(""),
):
    for date_value in (date_from, date_to):
        if date_value and not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", date_value):
            raise HTTPException(status_code=422, detail="Date filters must use YYYY-MM format.")
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=422, detail="The start month must not be after the end month.")
    path, sheet = _choose_dataset(_list_dataset_objects(), module_id, metric_name, metric_code, module_name, supabase_metric_name)
    payload = _read_dataset(
        path,
        _download_object(path),
        offset,
        limit,
        metric_name,
        metric_code,
        sheet,
        module_id,
        module_name,
        date_column,
        date_from,
        date_to,
    )
    payload["downloadUrl"] = (
        "/api/v1/datasets/download?"
        f"module_id={quote(module_id, safe='')}&metric_name={quote(metric_name, safe='')}&metric_code={quote(metric_code, safe='')}&module_name={quote(module_name, safe='')}&supabase_metric_name={quote(supabase_metric_name, safe='')}"
    )
    return payload


@router.get("/download")
def download_dataset(module_id: str = Query(...), metric_name: str = Query(...), metric_code: str = Query(""), module_name: str = Query(""), supabase_metric_name: str = Query("")):
    path, preferred_sheet = _choose_dataset(_list_dataset_objects(), module_id, metric_name, metric_code, module_name, supabase_metric_name)
    content = _download_object(path)
    if PurePosixPath(path).suffix.lower() == ".csv":
        file_name = PurePosixPath(path).name.replace('"', "")
        media_type = "text/csv; charset=utf-8"
    else:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = _select_worksheet(workbook, metric_name, metric_code, preferred_sheet, module_id, module_name)
            csv_output = io.StringIO(newline="")
            writer = csv.writer(csv_output)
            for row in sheet.iter_rows(values_only=True):
                if any(value is not None and str(value).strip() for value in row):
                    writer.writerow([_cell_value(value) for value in row])
            workbook.close()
            content = csv_output.getvalue().encode("utf-8-sig")
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Could not export the selected dataset worksheet: {exc}")
        sheet_part = re.sub(r"[^A-Za-z0-9._-]+", "_", sheet.title).strip("_.") or "dataset"
        file_name = f"{PurePosixPath(path).stem}_{sheet_part}.csv".replace('"', "")
        media_type = "text/csv; charset=utf-8"
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )
